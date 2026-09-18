"""Extract text (and, for PDFs, page images) from architecture/standards documents."""
from __future__ import annotations

import base64
import io
from dataclasses import dataclass, field
from pathlib import Path

import fitz  # pymupdf
import openpyxl
from docx import Document as DocxDocument


@dataclass
class Page:
    text: str
    image_b64: str | None = None  # PNG, base64-encoded


@dataclass
class ExtractedDocument:
    source: str
    pages: list[Page] = field(default_factory=list)

    @property
    def full_text(self) -> str:
        return "\n\n".join(p.text for p in self.pages if p.text)


def extract(path: str | Path, render_images: bool = True) -> ExtractedDocument:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf(path, render_images)
    if suffix in (".docx",):
        return _extract_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return _extract_xlsx(path)
    raise ValueError(f"Unsupported file type: {suffix} ({path})")


def _extract_pdf(path: Path, render_images: bool) -> ExtractedDocument:
    doc = fitz.open(path)
    pages = []
    for page in doc:
        text = page.get_text()
        image_b64 = None
        if render_images:
            pix = page.get_pixmap(dpi=150)
            image_b64 = base64.b64encode(pix.tobytes("png")).decode("ascii")
        pages.append(Page(text=text, image_b64=image_b64))
    doc.close()
    return ExtractedDocument(source=str(path), pages=pages)


def _extract_docx(path: Path) -> ExtractedDocument:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return ExtractedDocument(source=str(path), pages=[Page(text="\n".join(parts))])


def _extract_xlsx(path: Path) -> ExtractedDocument:
    wb = openpyxl.load_workbook(str(path), data_only=True)
    pages = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            if any(v is not None and str(v).strip() for v in row):
                rows.append(" | ".join("" if v is None else str(v) for v in row))
        if rows:
            pages.append(Page(text=f"[Sheet: {sheet.title}]\n" + "\n".join(rows)))
    return ExtractedDocument(source=str(path), pages=pages)
